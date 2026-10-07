#!/usr/bin/env python3
"""
ĐỒNG BỘ 2 CHIỀU TỰ ĐỘNG GIỮA HANA PM HUB (CLOUDFLARE KV) VÀ GOOGLE SHEETS / DRIVE
1. Công việc (Tasks): Google Sheet "04 Task Tracker" (ID: 1TxIBBRPTftXJP4oqmyDXidr-8mDFoybQZFpo6NBJsm8)
2. Mua sắm (Capex): Google Sheet "Setup Lễ tân" & "Setup Spa" (ID: 17abDmjThWZ-kQdW2cVPl2Kp8BfEz2v7trtebELIkD_s)
3. Hồ sơ Pháp lý (Legal): File Excel "BANG_THEO_DOI_TIEN_DO_PHAP_LY_VA_MO_CUA_HANA_WELLNESS.xlsx" (ID Drive: 1XpU-5goVpdFNgYGpV6wkVYznTA8KDsz5)
"""

import os
import sys
import json
import time
import argparse
import urllib.request
import urllib.parse
from datetime import datetime
import openpyxl

# Đường dẫn tệp token Google OAuth
TOKEN_PATH = "/Users/tungpv/.gemini/antigravity-ide/scratch/social-auto-agent/credentials/token.json"
CACHE_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), ".sync_cache.json")

# Google Sheets IDs
TASKS_SHEET_ID = "1TxIBBRPTftXJP4oqmyDXidr-8mDFoybQZFpo6NBJsm8"
CAPEX_SHEET_ID = "17abDmjThWZ-kQdW2cVPl2Kp8BfEz2v7trtebELIkD_s"

# Legal Excel path
LEGAL_EXCEL_PATH = "/Users/tungpv/Library/CloudStorage/GoogleDrive-phamtunghcm@gmail.com/My Drive/1 Hana Wellness Project/01 Kinh doanh Vận hành/02 SOP & Vận hành/BANG_THEO_DOI_TIEN_DO_PHAP_LY_VA_MO_CUA_HANA_WELLNESS.xlsx"
LEGAL_DRIVE_ID = "1XpU-5goVpdFNgYGpV6wkVYznTA8KDsz5"

# Cloudflare Pages API
KV_API_URL = "https://hana-pm-hub.pages.dev/api/data"
HEADERS = {"User-Agent": "Mozilla/5.0", "Content-Type": "application/json"}

def get_google_access_token():
    with open(TOKEN_PATH, "r", encoding="utf-8") as f:
        t_data = json.load(f)
    
    req = urllib.request.Request(
        "https://oauth2.googleapis.com/token",
        data=json.dumps({
            "client_id": t_data["client_id"],
            "client_secret": t_data["client_secret"],
            "refresh_token": t_data["refresh_token"],
            "grant_type": "refresh_token"
        }).encode(),
        headers={"Content-Type": "application/json"}
    )
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read().decode())["access_token"]

def fetch_kv_data():
    req = urllib.request.Request(KV_API_URL, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req) as resp:
        res = json.loads(resp.read().decode())
        if res.get("success") and res.get("data"):
            return res["data"]
    return None

def save_kv_data(data):
    req = urllib.request.Request(
        KV_API_URL,
        data=json.dumps(data).encode("utf-8"),
        headers=HEADERS,
        method="POST"
    )
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read().decode())

def load_cache():
    if os.path.exists(CACHE_PATH):
        try:
            with open(CACHE_PATH, "r", encoding="utf-8") as f:
                return json.load(f)
        except:
            pass
    return {"tasks": {}, "capex": {}, "legal": {}}

def save_cache(cache):
    with open(CACHE_PATH, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=2)

# ==========================================
# 1. TASKS SYNC (04 Task Tracker)
# ==========================================
def sync_tasks(token, kv_data, cache):
    print("\n--- [1] Kiểm tra & Đồng bộ Công việc (Tasks) ---")
    kv_tasks = {t["id"]: t for t in kv_data.get("tasks", []) if "id" in t}
    cached_tasks = cache.get("tasks", {})
    
    # Đọc Google Sheet
    rng = "04 Task Tracker!A5:J50"
    url = f"https://sheets.googleapis.com/v4/spreadsheets/{TASKS_SHEET_ID}/values/{urllib.parse.quote(rng)}"
    req = urllib.request.Request(url, headers={"Authorization": f"Bearer {token}"})
    with urllib.request.urlopen(req) as resp:
        rows = json.loads(resp.read().decode()).get("values", [])

    sheet_tasks = {}
    for i, r in enumerate(rows):
        row_num = 5 + i
        stt = r[1] if len(r) > 1 else ""
        title = r[3] if len(r) > 3 else ""
        status = r[7] if len(r) > 7 else ""
        try:
            tid = int(stt)
            sheet_tasks[tid] = {"row": row_num, "title": title, "status": status.strip()}
        except ValueError:
            continue

    # 1.1 Kiểm tra KV -> Sheet (User cập nhật trên web)
    updates_to_sheet = []
    kv_updated = False
    
    for tid, s_info in sheet_tasks.items():
        k_item = kv_tasks.get(tid)
        if not k_item:
            continue
        
        k_status = k_item.get("status", "").strip()
        s_status = s_info["status"]
        last_cached = cached_tasks.get(str(tid))

        # Nếu KV khác Sheet
        if k_status != s_status:
            # Ưu tiên bên nào thay đổi so với cache
            if last_cached is None or k_status != last_cached:
                # KV thay đổi -> Đẩy lên Sheet
                print(f"  [KV -> Sheet] Task #{tid} ('{s_info['title'][:30]}'): '{s_status}' -> '{k_status}'")
                updates_to_sheet.append({
                    "range": f"04 Task Tracker!H{s_info['row']}",
                    "values": [[k_status]]
                })
                cached_tasks[str(tid)] = k_status
            elif s_status != last_cached:
                # Sheet thay đổi -> Cập nhật vào KV
                print(f"  [Sheet -> KV] Task #{tid} ('{s_info['title'][:30]}'): '{k_status}' -> '{s_status}'")
                k_item["status"] = s_status
                cached_tasks[str(tid)] = s_status
                kv_updated = True
        else:
            cached_tasks[str(tid)] = s_status

    if updates_to_sheet:
        batch_body = json.dumps({
            "valueInputOption": "USER_ENTERED",
            "data": updates_to_sheet
        }).encode()
        req_b = urllib.request.Request(
            f"https://sheets.googleapis.com/v4/spreadsheets/{TASKS_SHEET_ID}/values:batchUpdate",
            data=batch_body,
            headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
            method="POST"
        )
        with urllib.request.urlopen(req_b) as resp:
            res_b = json.loads(resp.read().decode())
            print(f"  ✅ Đã cập nhật {res_b.get('totalUpdatedCells', len(updates_to_sheet))} ô trạng thái Task trên Google Sheet!")

    cache["tasks"] = cached_tasks
    return kv_updated

# ==========================================
# 2. CAPEX SYNC (Setup Lễ tân & Setup Spa)
# ==========================================
def sync_capex(token, kv_data, cache):
    print("\n--- [2] Kiểm tra & Đồng bộ Mua sắm (Capex) ---")
    kv_capex = {c["title"].strip(): c for c in kv_data.get("capex", []) if c.get("title")}
    cached_capex = cache.get("capex", {})
    updates_to_sheet = []
    kv_updated = False

    def process_subsheet(sheet_name, title_col, status_col, start_row, max_row):
        nonlocal kv_updated
        rng = f"{sheet_name}!A{start_row}:M{max_row}"
        url = f"https://sheets.googleapis.com/v4/spreadsheets/{CAPEX_SHEET_ID}/values/{urllib.parse.quote(rng)}"
        req = urllib.request.Request(url, headers={"Authorization": f"Bearer {token}"})
        with urllib.request.urlopen(req) as resp:
            rows = json.loads(resp.read().decode()).get("values", [])

        for i, r in enumerate(rows):
            row_num = start_row + i
            if len(r) > title_col:
                title = r[title_col].strip()
                s_status = r[status_col].strip() if len(r) > status_col else ""
                if not title or title not in kv_capex:
                    continue

                k_item = kv_capex[title]
                k_status = k_item.get("status", "").strip()
                last_cached = cached_capex.get(title)

                col_letter = chr(ord('A') + status_col)

                if k_status != s_status:
                    if last_cached is None or k_status != last_cached:
                        print(f"  [KV -> Sheet] {sheet_name} R{row_num} ('{title[:25]}'): '{s_status}' -> '{k_status}'")
                        updates_to_sheet.append({
                            "range": f"{sheet_name}!{col_letter}{row_num}",
                            "values": [[k_status]]
                        })
                        cached_capex[title] = k_status
                    elif s_status != last_cached:
                        print(f"  [Sheet -> KV] {sheet_name} R{row_num} ('{title[:25]}'): '{k_status}' -> '{s_status}'")
                        k_item["status"] = s_status
                        cached_capex[title] = s_status
                        kv_updated = True
                else:
                    cached_capex[title] = s_status

    process_subsheet("Setup Lễ tân", 5, 11, 8, 40)
    process_subsheet("Setup Spa", 4, 11, 8, 65)

    if updates_to_sheet:
        batch_body = json.dumps({
            "valueInputOption": "USER_ENTERED",
            "data": updates_to_sheet
        }).encode()
        req_b = urllib.request.Request(
            f"https://sheets.googleapis.com/v4/spreadsheets/{CAPEX_SHEET_ID}/values:batchUpdate",
            data=batch_body,
            headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
            method="POST"
        )
        with urllib.request.urlopen(req_b) as resp:
            res_b = json.loads(resp.read().decode())
            print(f"  ✅ Đã cập nhật {res_b.get('totalUpdatedCells', len(updates_to_sheet))} ô Capex trên Google Sheet!")

    cache["capex"] = cached_capex
    return kv_updated

# ==========================================
# 3. LEGAL SYNC (Excel file trên Google Drive qua Drive API v3)
# ==========================================
def sync_legal(token, kv_data, cache):
    print("\n--- [3] Kiểm tra & Đồng bộ Hồ sơ Pháp lý (Legal) ---")
    kv_legal = {l["id"]: l for l in kv_data.get("legal", []) if "id" in l}
    cached_legal = cache.get("legal", {})

    # Tải file trực tiếp từ Google Drive API v3
    try:
        url = f"https://www.googleapis.com/drive/v3/files/{LEGAL_DRIVE_ID}?alt=media"
        req = urllib.request.Request(url, headers={"Authorization": f"Bearer {token}"})
        with urllib.request.urlopen(req) as resp:
            data_bytes = resp.read()
    except Exception as e:
        print(f"  ⚠️ Không thể tải file Legal từ Google Drive API: {e}")
        return False

    import io
    wb = openpyxl.load_workbook(io.BytesIO(data_bytes))
    ws = wb["Chi tiết tiến độ"]
    
    excel_modified = False
    kv_updated = False

    for r in range(5, ws.max_row + 1):
        stt = ws.cell(row=r, column=1).value
        if stt is not None:
            try:
                stt_int = int(stt)
            except:
                continue
            title = str(ws.cell(row=r, column=3).value or "")
            e_status = str(ws.cell(row=r, column=11).value or "").strip()
            
            k_item = kv_legal.get(stt_int)
            if not k_item:
                continue
            
            k_status = str(k_item.get("status", "")).strip()
            last_cached = cached_legal.get(str(stt_int))

            if k_status != e_status:
                if last_cached is None or k_status != last_cached:
                    print(f"  [KV -> Excel] Legal #{stt_int} ('{title[:25]}'): '{e_status}' -> '{k_status}'")
                    ws.cell(row=r, column=11, value=k_status)
                    excel_modified = True
                    cached_legal[str(stt_int)] = k_status
                elif e_status != last_cached:
                    print(f"  [Excel -> KV] Legal #{stt_int} ('{title[:25]}'): '{k_status}' -> '{e_status}'")
                    k_item["status"] = e_status
                    cached_legal[str(stt_int)] = e_status
                    kv_updated = True
            else:
                cached_legal[str(stt_int)] = e_status

    if excel_modified:
        out_buf = io.BytesIO()
        wb.save(out_buf)
        up_bytes = out_buf.getvalue()

        # 1. Upload Google Drive API v3
        try:
            up_url = f"https://www.googleapis.com/upload/drive/v3/files/{LEGAL_DRIVE_ID}?uploadType=media"
            up_req = urllib.request.Request(
                up_url,
                data=up_bytes,
                headers={"Authorization": f"Bearer {token}", "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
                method="PATCH"
            )
            with urllib.request.urlopen(up_req) as up_resp:
                print("  ✅ Đã cập nhật file Excel thành công lên Google Drive API v3!")
        except Exception as e:
            print(f"  ⚠️ Lỗi tải file lên Drive API: {e}")

        # 2. Thử lưu bản sao local nếu có quyền ghi
        try:
            if os.path.exists(LEGAL_EXCEL_PATH):
                with open(LEGAL_EXCEL_PATH, "wb") as f_local:
                    f_local.write(up_bytes)
                print("  ✅ Đã đồng bộ file Excel vào ổ đĩa cục bộ!")
        except:
            pass

    cache["legal"] = cached_legal
    return kv_updated

# ==========================================
# MAIN RUNNER
# ==========================================
def run_full_sync():
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    print(f"\n==============================================")
    print(f"🚀 BẮT ĐẦU ĐỒNG BỘ 2 CHIỀU LÚC {now_str}")
    print(f"==============================================")
    
    token = get_google_access_token()
    kv_data = fetch_kv_data()
    if not kv_data:
        print("❌ Không lấy được dữ liệu từ Cloudflare KV!")
        return

    cache = load_cache()
    
    tasks_kv_changed = sync_tasks(token, kv_data, cache)
    capex_kv_changed = sync_capex(token, kv_data, cache)
    legal_kv_changed = sync_legal(token, kv_data, cache)

    save_cache(cache)

    if tasks_kv_changed or capex_kv_changed or legal_kv_changed:
        print("\n⚡ Phát hiện thay đổi từ Google Sheets/Excel -> Đang đẩy cập nhật lên Cloudflare KV...")
        res = save_kv_data(kv_data)
        print("Cloudflare KV response:", res)
    else:
        print("\n✨ Dữ liệu giữa PM Hub và Google Sheets / Excel đã đồng bộ 100%!")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Bidirectional sync between Hana PM Hub and Google Sheets")
    parser.add_argument("--daemon", action="store_true", help="Chạy chế độ daemon kiểm tra định kỳ mỗi 30 giây")
    parser.add_argument("--interval", type=int, default=30, help="Khoảng cách giữa các lần sync (giây)")
    args = parser.parse_args()

    if args.daemon:
        print(f"🔄 Bắt đầu chạy sync daemon (kiểm tra mỗi {args.interval} giây). Nhấn Ctrl+C để dừng.")
        while True:
            try:
                run_full_sync()
            except Exception as e:
                print(f"⚠️ Lỗi trong vòng lặp sync: {e}")
            time.sleep(args.interval)
    else:
        run_full_sync()
