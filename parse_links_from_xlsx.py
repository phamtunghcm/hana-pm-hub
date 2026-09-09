import openpyxl, json

xlsx_path = "/Users/tungpv/Library/CloudStorage/GoogleDrive-phamtunghcm@gmail.com/My Drive/1 Hana Wellness Project/01 Kinh doanh Vận hành/02 SOP & Vận hành/BANG_THEO_DOI_TIEN_DO_PHAP_LY_VA_MO_CUA_HANA_WELLNESS.xlsx"

wb = openpyxl.load_workbook(xlsx_path)
ws = wb["Chi tiết tiến độ"]

legal_files = []

# Row 4 is header: Col 1: STT, Col 2: Nhóm, Col 3: Hạng mục, Col 4: Nội dung, Col 5: Tài liệu liên quan (Click mở), Col 7: Cơ quan, Col 8: Phụ trách, Col 9: Bắt đầu, Col 10: Hạn chót, Col 11: Trạng thái, Col 12: Lưu ý
for r in range(5, ws.max_row + 1):
    stt = ws.cell(row=r, column=1).value
    group = ws.cell(row=r, column=2).value
    item = ws.cell(row=r, column=3).value
    content = ws.cell(row=r, column=4).value
    doc_cell = ws.cell(row=r, column=5).value
    legal_basis = ws.cell(row=r, column=6).value
    agency = ws.cell(row=r, column=7).value
    pic = ws.cell(row=r, column=8).value
    start = ws.cell(row=r, column=9).value
    deadline = ws.cell(row=r, column=10).value
    status = ws.cell(row=r, column=11).value
    note = ws.cell(row=r, column=12).value

    if stt is not None:
        # Extract hyperlink from formula if formula like =HYPERLINK("url", "label")
        url = ""
        label = str(doc_cell) if doc_cell else ""
        if doc_cell and isinstance(doc_cell, str) and "HYPERLINK(" in doc_cell:
            # extract url
            parts = doc_cell.split('"')
            if len(parts) >= 4:
                url = parts[1]
                label = parts[3]
        elif ws.cell(row=r, column=5).hyperlink:
            url = ws.cell(row=r, column=5).hyperlink.target
            
        legal_files.append({
            "stt": stt,
            "group": group,
            "title": item,
            "content": content,
            "fileLabel": label,
            "fileUrl": url,
            "legalBasis": legal_basis,
            "agency": agency,
            "pic": pic,
            "timeEstimate": f"{start} -> {deadline}",
            "status": status or "Chưa bắt đầu",
            "note": f"{content}. Căn cứ: {legal_basis}. Phụ trách: {pic}. {note}".strip()
        })

print(f"Parsed {len(legal_files)} items from Excel with exact Drive links:")
for it in legal_files[:5]:
    print(f"[{it['stt']}] {it['title']}")
    print(f"     Label: {it['fileLabel']}")
    print(f"     URL:   {it['fileUrl']}\n")
