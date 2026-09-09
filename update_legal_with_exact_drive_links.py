import openpyxl, urllib.request, json

xlsx_path = "/Users/tungpv/Library/CloudStorage/GoogleDrive-phamtunghcm@gmail.com/My Drive/1 Hana Wellness Project/01 Kinh doanh Vận hành/02 SOP & Vận hành/BANG_THEO_DOI_TIEN_DO_PHAP_LY_VA_MO_CUA_HANA_WELLNESS.xlsx"
wb = openpyxl.load_workbook(xlsx_path)
ws = wb["Chi tiết tiến độ"]

CLOUD_API = "https://hana-pm-hub.pages.dev/api/data"
HEADERS = {"User-Agent": "Mozilla/5.0"}

print("1. Đang đọc dữ liệu từ Cloudflare KV...")
req = urllib.request.Request(CLOUD_API, headers=HEADERS)
live_data = json.loads(urllib.request.urlopen(req).read().decode())["data"]
legal_map = {l.get("id"): l for l in live_data.get("legal", [])}

print("2. Đang đọc link chi tiết từng file từ bảng Excel theo dõi mới...")
new_legal = []
for r in range(5, ws.max_row + 1):
    stt_val = ws.cell(row=r, column=1).value
    if stt_val is None:
        continue
    try:
        stt = int(stt_val)
    except:
        continue
        
    group = ws.cell(row=r, column=2).value
    item = ws.cell(row=r, column=3).value
    content = ws.cell(row=r, column=4).value
    doc_cell = ws.cell(row=r, column=5).value
    legal_basis = ws.cell(row=r, column=6).value
    agency = ws.cell(row=r, column=7).value
    pic = ws.cell(row=r, column=8).value
    start = ws.cell(row=r, column=9).value
    deadline = ws.cell(row=r, column=10).value
    sheet_status = ws.cell(row=r, column=11).value
    note = ws.cell(row=r, column=12).value

    # Extract URL and label
    url = ""
    label = str(doc_cell) if doc_cell else ""
    if doc_cell and isinstance(doc_cell, str) and "HYPERLINK(" in doc_cell:
        parts = doc_cell.split('"')
        if len(parts) >= 4:
            url = parts[1]
            label = parts[3]
    elif ws.cell(row=r, column=5).hyperlink:
        url = ws.cell(row=r, column=5).hyperlink.target

    # Status: Giữ nguyên nếu đã sửa trên Web, ngược lại lấy từ sheet
    status = legal_map.get(stt, {}).get("status", sheet_status or "Chưa bắt đầu")

    new_legal.append({
        "id": stt,
        "group": group,
        "title": f"[{stt:02d}] {item}",
        "agency": agency or "Nội bộ",
        "timeEstimate": f"{start} -> {deadline}",
        "status": status,
        "note": f"{content}. Căn cứ: {legal_basis}. Phụ trách: {pic}. {note or ''}".strip(),
        "fileLabel": label,
        "fileLink": url, # Link mở trực tiếp từng file/thư mục trên Google Drive
        "sheetLink": "https://docs.google.com/spreadsheets/d/1XpU-5goVpdFNgYGpV6wkVYznTA8KDsz5/edit?gid=2104154183#gid=2104154183",
        "type": "legal"
    })

print(f"Đã xử lý {len(new_legal)} hồ sơ với link file Drive trực tiếp.")

# 3. Cập nhật lên Cloudflare KV
live_data["legal"] = new_legal

post_req = urllib.request.Request(
    CLOUD_API,
    headers={"Content-Type": "application/json", "User-Agent": "Mozilla/5.0"},
    data=json.dumps(live_data).encode("utf-8"),
    method="POST"
)
post_resp = urllib.request.urlopen(post_req)
print("Kết quả cập nhật Cloud KV:", post_resp.read().decode())

# 4. Lưu vào local file
with open("src/data/legal5.json", "w", encoding="utf-8") as f:
    json.dump(new_legal, f, ensure_ascii=False, indent=2)

print("Hoàn tất cập nhật link chi tiết từng file vào Database!")
